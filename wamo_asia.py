"""일본·대만 Lite 수집 엔진.

한국·미국처럼 큰 화면 파일(index.html/us.html)을 만들지 않고, 레이더에 필요한
결과만 작은 파일로 남긴다(저장소 용량 절약). 가격 이력은 매일 야후에서 다시 받는다.

- 종목 목록(주 1회 갱신, wamo_asia_universe.json)
  일본: 일본거래소(JPX) 공식 상장 목록의 TOPIX Core30·Large70·Mid400 = TOPIX 500.
        한글 종목명·상장주식수는 네이버 해외주식.
  대만: 대만증권거래소(TWSE)·장외(TPEx) 공식 자료의 상장주식수 × 종가 = 시가총액,
        300억 대만달러 이상.
- 시세: 야후 일봉 2년치(기존 한국·미국과 같은 함수 fetch_yahoo_history).
- 계산: 기존 calc_raw(RS 혼합값·Stage 2·이동평균)와 같은 백분위 → 한국·미국과 같은 기준.
- 결과: wamo_asia_markets.json(레이더 블록), wamo_chart_JP/TW.json(상세 차트),
        wamo_asia_status.json(갱신 상태). 레이더 화면은 wamo_publish_insights가 합친다.
"""
from __future__ import annotations

import argparse
import io
import json
import math
import os
import re
import time
import urllib.request
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parent
UNIVERSE_FILE = 'wamo_asia_universe.json'
MARKETS_FILE = 'wamo_asia_markets.json'
STATUS_FILE = 'wamo_asia_status.json'
UNIVERSE_DAYS = 7          # 종목 목록은 주 1회만 다시 만든다
UNIVERSE_MAX_AGE = 31      # 목록 갱신이 실패하면 이 기간까지는 예전 목록을 쓴다
TW_MIN_CAP = 3e10          # 300억 대만달러
JP_SIZES = ('TOPIX Core30', 'TOPIX Large70', 'TOPIX Mid400')
WORKERS = 8
UA = 'Mozilla/5.0 (WAMO asia lite)'

MARKETS = {
    'JP': {'label': '일본', 'bench': ('^N225', '닛케이 225'), 'currency': 'JPY',
           'universeLabel': 'TOPIX 500 (일본거래소 규모구분 Core30·Large70·Mid400)'},
    'TW': {'label': '대만', 'bench': ('^TWII', '대만 가권지수'), 'currency': 'TWD',
           'universeLabel': '대만 상장·장외 시가총액 300억 대만달러 이상'},
}

# 일본거래소 33업종 → 한국어
JP33 = {
    '水産・農林業': '수산·농림', '鉱業': '광업', '建設業': '건설', '食料品': '식품', '繊維製品': '섬유',
    'パルプ・紙': '펄프·종이', '化学': '화학', '医薬品': '의약품', '石油・石炭製品': '석유·석탄제품',
    'ゴム製品': '고무제품', 'ガラス・土石製品': '유리·요업', '鉄鋼': '철강', '非鉄金属': '비철금속',
    '金属製品': '금속제품', '機械': '기계', '電気機器': '전기기기', '輸送用機器': '수송기기(자동차·부품)',
    '精密機器': '정밀기기', 'その他製品': '기타제품', '電気・ガス業': '전기·가스', '陸運業': '육운',
    '海運業': '해운', '空運業': '항공운송', '倉庫・運輸関連業': '창고·운수', '情報・通信業': '정보·통신',
    '卸売業': '도매(종합상사)', '小売業': '소매', '銀行業': '은행', '証券、商品先物取引業': '증권',
    '保険業': '보험', 'その他金融業': '기타금융', '不動産業': '부동산', 'サービス業': '서비스',
}

# 대만 거래소 산업 코드(產業別) → 한국어
TW_INDUSTRY = {
    '01': '시멘트', '02': '식품', '03': '플라스틱', '04': '섬유', '05': '전기기계', '06': '전선·케이블',
    '08': '유리·도자기', '09': '제지', '10': '철강', '11': '고무', '12': '자동차', '14': '건설·건자재',
    '15': '해운', '16': '관광·외식', '17': '금융·보험', '18': '무역·백화점', '20': '기타', '21': '화학',
    '22': '바이오·의료', '23': '석유·가스·전력', '24': '반도체', '25': '컴퓨터·주변기기', '26': '광전자',
    '27': '통신·네트워크', '28': '전자부품', '29': '전자 유통', '30': '정보서비스', '31': '기타 전자',
    '32': '문화·콘텐츠', '33': '농업기술', '34': '전자상거래', '35': '친환경·에너지', '36': '디지털·클라우드',
    '37': '스포츠·레저', '38': '생활용품',
}


# ------------------------------------------------------------------ helpers
def http_get(url, timeout=40, retries=3):
    last = None
    for attempt in range(retries):
        try:
            req = urllib.request.Request(url, headers={'User-Agent': UA, 'Accept': 'application/json,text/html,*/*'})
            with urllib.request.urlopen(req, timeout=timeout) as r:
                return r.read()
        except Exception as exc:  # noqa: BLE001
            last = exc
            if getattr(exc, 'code', None) in (400, 401, 403, 404, 409, 410):
                break
            time.sleep(1.0 + attempt)
    raise RuntimeError(f'{url[:80]} · {type(last).__name__}: {str(last)[:120]}')


def http_json(url, **kw):
    return json.loads(http_get(url, **kw))


def num(value):
    if value is None or isinstance(value, bool):
        return None
    text = re.sub(r'[^0-9.\-]', '', str(value))
    try:
        out = float(text)
    except ValueError:
        return None
    return out if math.isfinite(out) else None


def read_json(path, default):
    try:
        return json.loads(Path(path).read_text(encoding='utf-8'))
    except (OSError, ValueError):
        return default


def write_json(path, value, compact=True):
    path = Path(path)
    tmp = path.with_suffix(path.suffix + '.tmp')
    tmp.write_text(json.dumps(value, ensure_ascii=False, allow_nan=False,
                              separators=(',', ':') if compact else None, indent=None if compact else 1) + '\n',
                   encoding='utf-8')
    tmp.replace(path)


def pick(row, *names, contains=()):
    """Field by exact name first, then by substring (field names differ between sources)."""
    for name in names:
        if name in row and str(row[name]).strip() not in ('', '-', '－'):
            return row[name]
    for key, value in row.items():
        if any(c in key for c in contains) and str(value).strip() not in ('', '-', '－'):
            return value
    return None


# ------------------------------------------------------------------ universe: Japan
def read_jpx(data):
    """JPX 상장종목 목록(xlsx/xls) → [{code, name, market, sector33, size}]."""
    if data[:2] == b'PK':
        from openpyxl import load_workbook
        sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
        rows = [list(r) for r in sheet.iter_rows(values_only=True)]
    else:
        import xlrd
        sheet = xlrd.open_workbook(file_contents=data).sheet_by_index(0)
        rows = [sheet.row_values(i) for i in range(sheet.nrows)]
    header = [str(h or '') for h in rows[0]]

    def col(*keys):
        for i, h in enumerate(header):
            if all(k in h for k in keys):
                return i
        raise ValueError(f'JPX 목록 열을 찾지 못함: {keys} · 머리글 {header}')
    ci, ni, mi, si, zi = col('コード'), col('銘柄名'), col('市場'), col('33業種区分'), col('規模区分')
    out = []
    for r in rows[1:]:
        code = r[ci]
        if isinstance(code, float) and code.is_integer():
            code = str(int(code))
        code = str(code or '').strip()
        if not re.fullmatch(r'[0-9][0-9A-Z]{3}', code):
            continue
        out.append({'code': code, 'name': str(r[ni] or '').strip(), 'market': str(r[mi] or ''),
                    'sector33': str(r[si] or '').strip(), 'size': str(r[zi] or '').strip()})
    return out


def jp_universe(get=http_get):
    page = get('https://www.jpx.co.jp/markets/statistics-equities/misc/01.html').decode('utf-8', 'replace')
    link = re.search(r'href="([^"]+data_j\.xlsx?)"', page)
    if not link:
        raise RuntimeError('JPX 상장종목 목록 파일 링크를 찾지 못함')
    url = link.group(1) if link.group(1).startswith('http') else 'https://www.jpx.co.jp' + link.group(1)
    listed = [r for r in read_jpx(get(url, timeout=90))
              if r['size'] in JP_SIZES and '内国株式' in r['market']]
    if len(listed) < 400:
        raise RuntimeError(f'TOPIX 500 종목이 너무 적음: {len(listed)}')

    def naver(item):
        try:
            d = json.loads(get(f"https://api.stock.naver.com/stock/{item['code']}.T/basic", timeout=20, retries=2))
            return item['code'], d.get('stockName'), num(d.get('countOfListedStock'))
        except Exception:  # noqa: BLE001 - Korean name is optional
            return item['code'], None, None
    with ThreadPoolExecutor(4) as pool:
        extra = {c: (n, s) for c, n, s in pool.map(naver, listed)}
    items = []
    for r in listed:
        kname, shares = extra.get(r['code'], (None, None))
        items.append({'code': r['code'], 'ticker': r['code'] + '.T', 'name': kname or r['name'], 'nameLocal': r['name'],
                      'industry': JP33.get(r['sector33'], r['sector33'] or '산업 미분류'), 'shares': shares,
                      'exchange': 'TSE', 'size': r['size']})
    named = sum(1 for i in items if i['shares'])
    return items, {'source': 'JPX 상장종목 목록 + 네이버 해외주식(한글명·상장주식수)', 'count': len(items),
                   'withShares': named}


# ------------------------------------------------------------------ universe: Taiwan
def tw_rows(info, quotes, exchange, suffix):
    close = {}
    for q in quotes:
        code = str(pick(q, 'Code', 'SecuritiesCompanyCode', '證券代號', '代號') or '').strip()
        price = num(pick(q, 'ClosingPrice', 'Close', '收盤價', '收盤'))
        if code and price:
            close[code] = price
    out = []
    for c in info:
        code = str(pick(c, '公司代號', 'SecuritiesCompanyCode', 'Code') or '').strip()
        if not re.fullmatch(r'[0-9]{4}', code):
            continue
        shares = num(pick(c, '已發行普通股數或TDR原股發行股數', contains=('已發行普通股數', 'IssueShares')))
        if not shares:
            capital = num(pick(c, '實收資本額', contains=('實收資本',)))
            shares = capital / 10 if capital else None   # 액면 10대만달러
        price = close.get(code)
        cap = shares * price if shares and price else None
        zh = str(pick(c, '公司簡稱', 'CompanyAbbreviation', contains=('簡稱',)) or code).strip()
        en = str(pick(c, '英文簡稱', contains=('英文簡稱', 'EnglishAbbreviation')) or '').strip()
        ind = str(pick(c, '產業別', 'SecuritiesIndustryCode', contains=('產業別',)) or '').strip().zfill(2)
        out.append({'code': code, 'ticker': code + suffix, 'name': f'{en} ({zh})' if en and en != zh else zh,
                    'nameLocal': zh, 'industry': TW_INDUSTRY.get(ind, '기타'), 'shares': shares,
                    'exchange': exchange, 'cap': cap})
    return out


def tw_universe(get_json=http_json):
    notes = []
    rows = tw_rows(get_json('https://openapi.twse.com.tw/v1/opendata/t187ap03_L'),
                   get_json('https://openapi.twse.com.tw/v1/exchangeReport/STOCK_DAY_ALL'), 'TWSE', '.TW')
    try:
        otc = tw_rows(get_json('https://www.tpex.org.tw/openapi/v1/mopsfin_t187ap03_O'),
                      get_json('https://www.tpex.org.tw/openapi/v1/tpex_mainboard_quotes'), 'TPEx', '.TWO')
        priced = sum(1 for r in otc if r['cap'])
        notes.append(f'장외(TPEx) {len(otc)}종목 중 시가총액 계산 {priced}종목')
        rows += otc
    except Exception as exc:  # noqa: BLE001 - listed board alone still works
        notes.append(f'장외(TPEx) 자료 실패 · 상장(TWSE)만 사용: {str(exc)[:120]}')
    items = sorted((r for r in rows if r['cap'] and r['cap'] >= TW_MIN_CAP), key=lambda r: -r['cap'])
    if len(items) < 120:
        raise RuntimeError(f'대만 시가총액 기준 종목이 너무 적음: {len(items)}')
    for r in items:
        r.pop('cap', None)
    return items, {'source': '대만증권거래소·장외시장 공식 공개자료(OpenAPI)', 'count': len(items), 'notes': notes}


def load_universe(market, root=ROOT, today=None, builders=None):
    today = today or date.today()
    store = read_json(root / UNIVERSE_FILE, {})
    old = store.get(market) or {}
    try:
        age = (today - date.fromisoformat(old.get('builtAt', '')[:10])).days
    except ValueError:
        age = None
    if age is not None and age < UNIVERSE_DAYS and old.get('items'):
        return old['items'], old.get('info') or {}, False
    builders = builders or {'JP': jp_universe, 'TW': tw_universe}
    try:
        items, info = builders[market]()
        store[market] = {'builtAt': today.isoformat(), 'items': items, 'info': info}
        write_json(root / UNIVERSE_FILE, store, compact=False)
        return items, info, True
    except Exception as exc:  # noqa: BLE001
        if old.get('items') and age is not None and age <= UNIVERSE_MAX_AGE:
            info = dict(old.get('info') or {}, warning=f'종목 목록 새로 만들기 실패 · {old.get("builtAt")} 목록 사용: {str(exc)[:120]}')
            return old['items'], info, False
        raise


# ------------------------------------------------------------------ prices and payload
def fetch_histories(items, expected, fetch=None):
    import wamo_update_business_dart as core
    fetch = fetch or (lambda t: core.fetch_yahoo_history(t, years=2, expected_date=expected)[0])

    def one(item):
        try:
            return item, fetch(item['ticker']), None
        except Exception as exc:  # noqa: BLE001
            return item, None, str(exc)[:160]
    with ThreadPoolExecutor(WORKERS) as pool:
        return list(pool.map(one, items))


def build_payload(market, items, results, expected):
    import wamo_update_business_dart as core
    stocks, errors = [], []
    for item, rows, err in results:
        if not rows:
            errors.append(f"{item['ticker']}: {err}")
            continue
        raw = core.calc_raw({}, rows)
        cap = item['shares'] * raw['close'] if item.get('shares') else None
        stocks.append({**raw, 'stock_code': item['code'], 'ticker': item['ticker'], 'name': item['name'],
                       'nameLocal': item.get('nameLocal'), 'industry': item['industry'], 'exchange': item['exchange'],
                       'market_cap_local': cap, 'dataSource': 'Yahoo Finance', 'priceProvider': 'Yahoo Finance'})
    values = [s['rsBlend'] for s in stocks]
    for s in stocks:
        s['rsPercentile'] = round(core.percentile(values, s['rsBlend']), 1)
    warnings = []
    if errors:
        warnings.append(f'가격을 받지 못한 {len(errors)}종목은 제외했습니다.')
    return {'meta': {'asOf': expected, 'market': market, 'source': f"{MARKETS[market]['universeLabel']} · 야후 일봉",
                     'warnings': warnings, 'currency': MARKETS[market]['currency']},
            'stocks': stocks}, errors


def chart_file(market, payload, bench, asof):
    from wamo_trend import chart_payload, clean_rows
    universe = [{'symbol': s['stock_code'], 'rows': clean_rows(s['history'])} for s in payload['stocks']]
    charts = chart_payload(universe, bench or [], asof, market)
    charts['benchLabel'] = MARKETS[market]['bench'][1]
    return charts


def trend_universe(market, payload):
    """Same per-stock shape wamo_trend builds for Korea/US (no earnings data for Japan/Taiwan)."""
    from wamo_trend import clean_rows
    out = []
    for s in payload['stocks']:
        rows = clean_rows(s.get('history'))
        if not rows:
            continue
        out.append({'id': f"{market}:{s['stock_code']}", 'symbol': s['stock_code'], 'ticker': s['ticker'],
                    'name': s['name'], 'industry': s.get('industry') or '산업 미분류',
                    'marketCap': s.get('market_cap_local'), 'rows': rows, 'closes': [r['close'] for r in rows],
                    'rsPercentile': s.get('rsPercentile'), 'vcp': None,
                    'growth': {'ok': None, 'text': '실적 자료 없음'}})
    return out


def build_trend(market, payload, asof, root, now, index_fetcher=None):
    """Trend signal, FTD and leaders for Japan/Taiwan, written into the shared wamo_trend.json."""
    import wamo_trend as trend
    log = trend.read_json(root / trend.SIGNAL_LOG, {'schemaVersion': 1, 'markets': {}})
    result = trend.read_json(root / trend.OUTPUT, {'schemaVersion': 1, 'markets': {}})
    prev = (result.get('markets') or {}).get(market)
    try:
        block, _, charts = trend.build_block(root, market, asof, trend_universe(market, payload), log, {},
                                             True, index_fetcher or trend.fetch_index,
                                             allow_synthetic=not trend.has_real_index(prev))
    except trend.IndexDelayed as exc:
        # Keep the last real-index verdict; the scheduled retries rebuild once the bar arrives.
        result['markets'][market] = trend.mark_kept(prev, '지수 자료 지연', now)
        trend.write_json(root / trend.OUTPUT, result, compact=True)
        print(f'{market} 추세 신호: {exc}', flush=True)
        return result['markets'][market], None
    block['builtAt'] = now.isoformat()
    result.setdefault('markets', {})[market] = block
    result['generatedAt'] = now.isoformat()
    trend.write_json(root / trend.OUTPUT, result, compact=True)
    trend.write_json(root / trend.SIGNAL_LOG, log, compact=True)
    print(f"{market} 추세 신호: 날씨 {block['weather']['label']} · FTD {block['ftd']['indexes'][0].get('phase')} · "
          f"선도주 {len((block.get('leadership') or {}).get('leaders') or [])}", flush=True)
    return block, charts


# ------------------------------------------------------------------ main
def run(markets=('JP', 'TW'), root=ROOT, now=None, force=False, fetch=None, index_fetcher=None, builders=None):
    import wamo_insights as insights
    from wamo_runtime import expected_session
    now = now or datetime.now(timezone.utc)
    # The radar file already holds the last published JP/TW blocks; MARKETS_FILE is only a
    # hand-off to wamo_publish_insights within this run (not committed, saves repo space).
    stored = read_json(root / MARKETS_FILE, None) or {'schemaVersion': 1, 'markets': {
        m: b for m, b in (read_json(root / 'wamo_insights.json', {}).get('markets') or {}).items() if m in MARKETS}}
    status = read_json(root / STATUS_FILE, {})
    changed, failed = [], []
    for market in markets:
        expected = expected_session(market, now)
        prev = (stored.get('markets') or {}).get(market)
        entry = dict(status.get(market) or {}, attemptedAt=now.isoformat(), expectedSession=expected,
                     runUrl=(f"https://github.com/{os.environ['GITHUB_REPOSITORY']}/actions/runs/{os.environ['GITHUB_RUN_ID']}"
                             if os.getenv('GITHUB_RUN_ID') else None))
        # Yahoo often posts the Asian index's volume/low a few hours after the close. Keep the
        # scheduled retries running until that day's index bar is complete (FTD can judge it).
        trend_block = (read_json(root / 'wamo_trend.json', {}).get('markets') or {}).get(market) or {}
        ftd0 = ((trend_block.get('ftd') or {}).get('indexes') or [{}])[0]
        index_done = trend_block.get('asOf') == expected and ftd0.get('status') == 'OK'
        if not force and prev and prev.get('asOf') == expected and prev.get('status') == 'CURRENT' and index_done:
            print(f'{market} {expected} 이미 반영됨 · 건너뜀', flush=True)
            continue
        try:
            items, info, rebuilt = load_universe(market, root, now.date(), builders)
            t0 = time.time()
            results = fetch_histories(items, expected, fetch)
            payload, errors = build_payload(market, items, results, expected)
            if info.get('warning'):
                payload['meta']['warnings'].append(info['warning'])
            ok = len(payload['stocks'])
            print(f'{market} {expected}: 종목 {len(items)} · 가격 {ok} · 실패 {len(errors)} · {time.time() - t0:.0f}초'
                  + (' · 종목 목록 새로 만듦' if rebuilt else ''), flush=True)
            if ok < len(items) * 0.9:
                raise RuntimeError(f'가격 확보 {ok}/{len(items)} — 90% 미만이라 이번 결과를 쓰지 않음 · 예: {errors[:3]}')
            block = insights.build_market(market, payload, now, prev, None)
            block['universe'] = {'label': MARKETS[market]['universeLabel'], **{k: v for k, v in info.items() if k != 'warning'}}
            stored.setdefault('markets', {})[market] = block
            charts = None
            try:
                _, charts = build_trend(market, payload, expected, root, now, index_fetcher)
                changed += ['wamo_trend.json', 'wamo_signal_log.json', 'wamo_ftd_history.json']
            except Exception as exc:  # noqa: BLE001 - the stock scan still publishes; trend keeps its last block
                print(f'{market} 추세·FTD·선도주 계산 실패 · 이전 결과 유지: {type(exc).__name__}: {exc}', flush=True)
                payload['meta']['warnings'].append('추세·FTD·선도주 계산에 실패해 이전 결과를 보여 줍니다.')
            if charts is None:
                bench = None
                try:
                    from wamo_trend import fetch_index
                    bench = [r for r in (index_fetcher or fetch_index)(MARKETS[market]['bench'][0]) if r['date'] <= expected]
                except Exception as exc:  # noqa: BLE001 - chart works without the index line
                    print(f'{market} 지수 이력 실패 · 차트의 지수 비교선 생략: {exc}', flush=True)
                charts = chart_file(market, payload, bench, expected)
            write_json(root / f'wamo_chart_{market}.json', charts)
            entry.update(status='PASS' if block.get('status') == 'CURRENT' else 'WARNING', lastSuccessAt=now.isoformat(),
                         asOf=block.get('asOf'), count=ok, universe=len(items), failed=len(errors), error=None)
            changed += [f'wamo_chart_{market}.json'] + ([UNIVERSE_FILE] if rebuilt else [])
        except Exception as exc:  # noqa: BLE001 - keep the previous block, record the failure
            print(f'{market} 갱신 실패 · 이전 결과 유지: {type(exc).__name__}: {exc}', flush=True)
            entry.update(status='FAILED', error=str(exc)[:300])
            failed.append(market)
        status[market] = entry
        changed.append(STATUS_FILE)
    stored['generatedAt'] = now.isoformat()
    write_json(root / MARKETS_FILE, stored)
    write_json(root / STATUS_FILE, status, compact=False)
    return sorted(set(changed)), status, failed


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--markets', default='JP,TW')
    parser.add_argument('--force', action='store_true')
    parser.add_argument('--publish', action='store_true', help='레이더 화면을 다시 만들고 GitHub에 올림')
    args = parser.parse_args()
    names, status, failed = run(tuple(m.strip().upper() for m in args.markets.split(',') if m.strip()), force=args.force)
    if names and args.publish:
        from wamo_publish_insights import publish_snapshot
        publish_snapshot(ROOT)
        from wamo_run import publish
        publish(names + ['radar.html', 'wamo_insights.json'])
    if failed:
        raise SystemExit(1)
